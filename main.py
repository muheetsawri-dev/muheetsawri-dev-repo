"""
FaceTrack Pro — Smart Attendance & Face Recognition Suite (Tabbed Edition)

This is a restructure of the original single-screen app into 6 tabs using
CustomTkinter, on top of the SAME backend (Excel storage, camera capture,
LBPH training/recognition) as before.

FOLDER LAYOUT (unchanged — matches what you already have in VS Code):
    your-project/
      main.py                              <- this file
      haarcascade_frontalface_default.xml  <- must sit next to this file
      LICENSE
      README
      StudentDetails/
        StudentDetails.xlsx
      TrainingImage/
        <name>.<serial>.<id>.<n>.jpg
      TrainingImageLabel/
        Trainner.yml
        psd.txt
        settings.json           <- NEW: stores Settings-tab values
      Attendance/
        Attendance_<dd-mm-yyyy>.xlsx

Every path below is computed relative to THIS file's location
(BASE_DIR = os.path.dirname(os.path.abspath(__file__))), so as long as
main.py stays in the project root next to the cascade file, nothing needs
to change when you move the folder or open it on another machine.

Requires:
    pip install customtkinter opencv-contrib-python pillow openpyxl numpy
"""

import customtkinter as ctk
import tkinter as tk
from tkinter import ttk
import cv2
import os
import json
import hashlib
import threading
import platform
import subprocess
import numpy as np
from PIL import Image
import datetime
import time
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook

# ----------------------------------------------------------------------
# Paths — all relative to this file, matches your existing folder layout
# ----------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STUDENT_DETAILS_DIR = os.path.join(BASE_DIR, "StudentDetails")
STUDENT_DETAILS_XLSX = os.path.join(STUDENT_DETAILS_DIR, "StudentDetails.xlsx")
TRAINING_IMAGE_DIR = os.path.join(BASE_DIR, "TrainingImage")
TRAINING_LABEL_DIR = os.path.join(BASE_DIR, "TrainingImageLabel")
TRAINER_FILE = os.path.join(TRAINING_LABEL_DIR, "Trainner.yml")
PASSWORD_FILE = os.path.join(TRAINING_LABEL_DIR, "psd.txt")
SETTINGS_FILE = os.path.join(TRAINING_LABEL_DIR, "settings.json")
ATTENDANCE_DIR = os.path.join(BASE_DIR, "Attendance")
HAARCASCADE_PATH = os.path.join(BASE_DIR, "haarcascade_frontalface_default.xml")
HAARCASCADE_RAW_URL = ("https://raw.githubusercontent.com/opencv/opencv/master/data/"
                        "haarcascades/haarcascade_frontalface_default.xml")

MAX_SAMPLES = 60
CAM_WIDTH, CAM_HEIGHT = 640, 480
DETECT_EVERY_N_FRAMES = 2
CAM_WARMUP_ATTEMPTS = 15
MIN_SAMPLES_WARN = 20

for _d in (STUDENT_DETAILS_DIR, TRAINING_IMAGE_DIR, TRAINING_LABEL_DIR, ATTENDANCE_DIR):
    os.makedirs(_d, exist_ok=True)

busy_lock = threading.Lock()

# ----------------------------------------------------------------------
# Settings — everything that used to be a hardcoded constant now lives
# in TrainingImageLabel/settings.json and is editable from the Settings tab
# ----------------------------------------------------------------------
DEFAULT_SETTINGS = {
    "class_start_time": "09:00",   # HH:MM, 24h
    "confidence_threshold": 50,     # lower = stricter match
    "camera_index": 0,
    "sound_on": True,
}


def load_settings():
    if os.path.isfile(SETTINGS_FILE):
        try:
            with open(SETTINGS_FILE, "r") as f:
                data = json.load(f)
            merged = dict(DEFAULT_SETTINGS)
            merged.update(data)
            return merged
        except Exception:
            pass
    return dict(DEFAULT_SETTINGS)


def save_settings(new_settings):
    os.makedirs(TRAINING_LABEL_DIR, exist_ok=True)
    with open(SETTINGS_FILE, "w") as f:
        json.dump(new_settings, f, indent=2)


SETTINGS = load_settings()


def class_start_time_obj():
    h, m = SETTINGS["class_start_time"].split(":")
    return datetime.time(int(h), int(m), 0)


# ----------------------------------------------------------------------
# Palette — light "school management" theme: indigo sidebar, white cards
# ----------------------------------------------------------------------
BG_APP = "#f5f5fc"        # page background (soft lavender-white)
BG_PANEL = "#ffffff"      # card / panel background
BORDER = "#e7e7f3"
TEXT_LIGHT = "#2b2b3d"    # primary text (dark, sits on light panels)
TEXT_MUTED = "#8b8ba3"
GOLD = "#6c5ce7"          # brand / avatar accent (was gold, now indigo)
TEAL = "#5b4fe9"          # primary actions (indigo)
GREEN = "#22c55e"         # on-time / success
AMBER = "#f59e0b"         # late / warning
RED = "#ef4444"           # danger / delete

SIDEBAR_BG = "#5b4fe9"        # indigo sidebar
SIDEBAR_ACTIVE = "#7669f0"    # active/hover nav item
SIDEBAR_TEXT = "#e9e7fd"
SIDEBAR_TEXT_MUTED = "#c3bdf7"

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

FONT_MONO = ("Consolas", 11)


# ----------------------------------------------------------------------
# Small helpers (unchanged behavior from the original app)
# ----------------------------------------------------------------------
def assure_path_exists(path):
    directory = path if os.path.isdir(path) or path.endswith(os.sep) else os.path.dirname(path)
    if directory and not os.path.exists(directory):
        os.makedirs(directory)


def hash_password(raw):
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def open_camera():
    """Open the configured webcam using DirectShow on Windows first."""
    idx = SETTINGS.get("camera_index", 0)
    try:
        idx = int(idx)
    except (TypeError, ValueError):
        idx = 0

    if platform.system() == "Windows":
        cam = cv2.VideoCapture(idx, cv2.CAP_DSHOW)
        if cam is not None and cam.isOpened():
            return cam
        if cam is not None:
            cam.release()

    cam = cv2.VideoCapture(idx)
    if cam is not None and cam.isOpened():
        return cam
    if cam is not None:
        cam.release()
    return cv2.VideoCapture()


def warm_up_camera(cam, timeout_seconds=4.0):
    """Wait until the webcam returns a valid frame."""
    deadline = time.time() + timeout_seconds
    delay = 0.05
    while time.time() < deadline:
        ret, frame = cam.read()
        if ret and frame is not None and frame.size > 0:
            return True
        time.sleep(delay)
        delay = min(delay * 1.4, 0.3)
    return False

def cascade_is_valid(path):
    if not os.path.isfile(path):
        return False, f"File not found at:\n{path}"
    if os.path.getsize(path) < 10_000:
        return False, ("The file is far too small to be a real cascade file "
                        "(it may have been saved as an HTML page instead of the raw XML).")
    try:
        clf = cv2.CascadeClassifier(path)
    except Exception as exc:
        return False, f"OpenCV could not parse the file: {exc}"
    if clf.empty():
        return False, ("OpenCV loaded the file but it is not a valid cascade "
                        "(check that it starts with <?xml ... in a text editor).")
    return True, ""


def compute_status(dt_obj):
    return "Late" if dt_obj.time() > class_start_time_obj() else "On Time"


def compute_status_from_timestr(tstamp):
    try:
        t = datetime.datetime.strptime(str(tstamp), "%H:%M:%S").time()
        return "Late" if t > class_start_time_obj() else "On Time"
    except Exception:
        return "On Time"


def ui(func, *args):
    """Schedule a UI update safely from a background thread."""
    def _run():
        try:
            func(*args)
        except tk.TclError:
            pass
    try:
        app.after(0, _run)
    except tk.TclError:
        pass


def beep():
    if SETTINGS.get("sound_on", True):
        try:
            app.bell()
        except tk.TclError:
            pass


# ----------------------------------------------------------------------
# Excel storage helpers (unchanged logic from the original app)
# ----------------------------------------------------------------------
def _ensure_student_workbook():
    assure_path_exists(STUDENT_DETAILS_DIR)
    if not os.path.isfile(STUDENT_DETAILS_XLSX):
        wb = Workbook()
        ws = wb.active
        ws.title = "StudentDetails"
        ws.append(['SERIAL NO.', 'ID', 'NAME'])
        wb.save(STUDENT_DETAILS_XLSX)
        wb.close()


def read_student_details():
    """Returns (next_serial, {serial: (id, name)})."""
    _ensure_student_workbook()
    wb = load_workbook(STUDENT_DETAILS_XLSX, read_only=True)
    ws = wb.active
    data = {}
    max_serial = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        serial, sid, name = row[0], row[1], row[2]
        try:
            serial = int(serial)
        except (TypeError, ValueError):
            continue
        data[serial] = (str(sid), str(name))
        max_serial = max(max_serial, serial)
    wb.close()
    return max_serial + 1, data


def append_student_detail(serial, Id, name):
    _ensure_student_workbook()
    wb = load_workbook(STUDENT_DETAILS_XLSX)
    ws = wb.active
    ws.append([serial, Id, name])
    wb.save(STUDENT_DETAILS_XLSX)
    wb.close()


def remove_student_detail(serial):
    _ensure_student_workbook()
    wb = load_workbook(STUDENT_DETAILS_XLSX)
    ws = wb.active
    rows_to_delete = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row and row[0] is not None:
            try:
                if int(row[0]) == serial:
                    rows_to_delete.append(i)
            except (TypeError, ValueError):
                continue
    for i in reversed(rows_to_delete):
        ws.delete_rows(i)
    wb.save(STUDENT_DETAILS_XLSX)
    wb.close()


def delete_training_images(serial):
    if not os.path.isdir(TRAINING_IMAGE_DIR):
        return 0
    removed = 0
    for fname in os.listdir(TRAINING_IMAGE_DIR):
        parts = fname.split(".")
        if len(parts) >= 3:
            try:
                if int(parts[1]) == serial:
                    os.remove(os.path.join(TRAINING_IMAGE_DIR, fname))
                    removed += 1
            except (ValueError, OSError):
                continue
    return removed


def photo_count_for_serial(serial):
    if not os.path.isdir(TRAINING_IMAGE_DIR):
        return 0
    n = 0
    for fname in os.listdir(TRAINING_IMAGE_DIR):
        parts = fname.split(".")
        if len(parts) >= 3:
            try:
                if int(parts[1]) == serial:
                    n += 1
            except ValueError:
                continue
    return n


def append_attendance_rows(date_str, rows):
    assure_path_exists(ATTENDANCE_DIR)
    path = os.path.join(ATTENDANCE_DIR, f"Attendance_{date_str}.xlsx")
    if os.path.isfile(path):
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.append(['Id', 'Name', 'Date', 'In-Time', 'Status'])
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()
    return path


def load_today_attendance():
    date_str = datetime.datetime.now().strftime('%d-%m-%Y')
    return load_attendance_for_date(date_str), date_str


def load_attendance_for_date(date_str):
    """Returns a list of (id, name, date, time, status) tuples, newest first."""
    path = os.path.join(ATTENDANCE_DIR, f"Attendance_{date_str}.xlsx")
    records = []
    if os.path.isfile(path):
        wb = load_workbook(path, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            vals = list(row) + [None] * (5 - len(row))
            sid, name_, date_val, tstamp, status = vals[:5]
            if not status:
                status = compute_status_from_timestr(tstamp)
            records.append((str(sid), str(name_), str(date_val), str(tstamp), str(status)))
        wb.close()
    records.reverse()
    return records


def load_attendance_range(start_date, end_date):
    """start_date/end_date: datetime.date. Returns list of all records in range."""
    all_records = []
    d = start_date
    while d <= end_date:
        date_str = d.strftime('%d-%m-%Y')
        all_records.extend(load_attendance_for_date(date_str))
        d += datetime.timedelta(days=1)
    return all_records


# ----------------------------------------------------------------------
# Live state shared across tabs
# ----------------------------------------------------------------------
attendance_rows_cache = []   # (id, name, date, time, status), newest first
today_date_str = ""
app = None                   # set once the CTk root exists


# ============================================================================
# App
# ============================================================================
class FaceTrackApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("FaceTrack Pro — Attendance Management System")
        self.geometry("1280x820")
        self.minsize(1080, 700)
        self.configure(fg_color=BG_APP)

        self._build_tabs()
        self._build_statusbar()

        self.refresh_all()

    # ---- sidebar + page switching -----------------------------------------
    NAV_ITEMS = (
        ("Dashboard", "\u25a6"),
        ("Take Attendance", "\u25c9"),
        ("Students", "\u25a4"),
        ("Register", "\u2295"),
        ("Reports", "\u25b2"),
        ("Settings", "\u2699"),
    )

    def _build_tabs(self):
        # Outer row: sidebar on the left, content area on the right.
        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(fill="both", expand=True)
        body.grid_columnconfigure(1, weight=1)
        body.grid_rowconfigure(0, weight=1)

        sidebar = ctk.CTkFrame(body, fg_color=SIDEBAR_BG, corner_radius=0, width=220)
        sidebar.grid(row=0, column=0, sticky="nsw")
        sidebar.grid_propagate(False)

        brand = ctk.CTkFrame(sidebar, fg_color="transparent")
        brand.pack(fill="x", padx=20, pady=(24, 20))
        ctk.CTkLabel(brand, text="FT", width=36, height=36, corner_radius=8,
                     fg_color="#ffffff", text_color=SIDEBAR_BG,
                     font=("Consolas", 13, "bold")).pack(side="left")
        title_block = ctk.CTkFrame(brand, fg_color="transparent")
        title_block.pack(side="left", padx=(10, 0))
        ctk.CTkLabel(title_block, text="FaceTrack Pro", text_color="#ffffff",
                     font=("Segoe UI", 13, "bold")).pack(anchor="w")
        ctk.CTkLabel(title_block, text="Attendance suite", text_color=SIDEBAR_TEXT_MUTED,
                     font=("Segoe UI", 9)).pack(anchor="w")

        self._nav_buttons = {}
        nav_wrap = ctk.CTkFrame(sidebar, fg_color="transparent")
        nav_wrap.pack(fill="x", padx=12)
        for name, icon in self.NAV_ITEMS:
            btn = ctk.CTkButton(
                nav_wrap, text=f"  {icon}   {name}", anchor="w", corner_radius=8,
                fg_color="transparent", hover_color=SIDEBAR_ACTIVE,
                text_color=SIDEBAR_TEXT, font=("Segoe UI", 12), height=38,
                command=lambda n=name: self._select_page(n),
            )
            btn.pack(fill="x", pady=3)
            self._nav_buttons[name] = btn

        ctk.CTkFrame(sidebar, fg_color=SIDEBAR_TEXT_MUTED, height=1).pack(fill="x", padx=20, pady=16)
        ok, _reason = cascade_is_valid(HAARCASCADE_PATH)
        cascade_txt = "\u25cf Cascade file OK" if ok else "\u25cf Cascade file invalid"
        ctk.CTkLabel(sidebar, text=cascade_txt, text_color=(SIDEBAR_TEXT_MUTED if ok else "#ffb4b4"),
                     font=("Segoe UI", 9)).pack(anchor="w", padx=20)

        # Content area
        content = ctk.CTkFrame(body, fg_color="transparent")
        content.grid(row=0, column=1, sticky="nsew")
        self._build_header_in(content)

        self._pages = {}
        pages_host = ctk.CTkFrame(content, fg_color="transparent")
        pages_host.pack(fill="both", expand=True, padx=24, pady=(0, 8))
        for name, _icon in self.NAV_ITEMS:
            page = ctk.CTkFrame(pages_host, fg_color=BG_PANEL, corner_radius=12,
                                 border_color=BORDER, border_width=1)
            self._pages[name] = page

        self._build_dashboard_tab(self._pages["Dashboard"])
        self._build_attendance_tab(self._pages["Take Attendance"])
        self._build_students_tab(self._pages["Students"])
        self._build_register_tab(self._pages["Register"])
        self._build_reports_tab(self._pages["Reports"])
        self._build_settings_tab(self._pages["Settings"])

        self._select_page("Dashboard")

    def _build_header_in(self, content):
        header = ctk.CTkFrame(content, fg_color="transparent")
        header.pack(fill="x", padx=24, pady=(20, 12))
        self.header_title = ctk.CTkLabel(header, text="Dashboard", text_color=TEXT_LIGHT,
                                          font=("Segoe UI", 20, "bold"))
        self.header_title.pack(side="left")
        self.clock_label = ctk.CTkLabel(header, text="", text_color=TEXT_MUTED, font=("Consolas", 13))
        self.clock_label.pack(side="right")
        self._tick_clock()

    def _tick_clock(self):
        self.clock_label.configure(text=time.strftime("%H:%M:%S — %d %b %Y"))
        self.after(1000, self._tick_clock)

    def _select_page(self, name):
        for n, btn in self._nav_buttons.items():
            btn.configure(fg_color=SIDEBAR_ACTIVE if n == name else "transparent")
        for n, page in self._pages.items():
            if n == name:
                page.pack(fill="both", expand=True)
            else:
                page.pack_forget()
        self.header_title.configure(text=name)

    # ---- status bar ---------------------------------------------------
    def _build_statusbar(self):
        bar = ctk.CTkFrame(self, fg_color="#eceafb", height=30, corner_radius=0)
        bar.pack(fill="x", side="bottom")
        self.status_label = ctk.CTkLabel(bar, text="● Ready", text_color=GREEN, font=("Segoe UI", 10))
        self.status_label.pack(side="left", padx=16, pady=4)

        ok, _reason = cascade_is_valid(HAARCASCADE_PATH)
        cascade_txt = "Cascade file OK" if ok else "Cascade file invalid — see Register tab"
        ctk.CTkLabel(bar, text=cascade_txt, text_color=(TEXT_MUTED if ok else RED),
                     font=("Segoe UI", 10)).pack(side="right", padx=16, pady=4)

    def set_status(self, text, busy=False):
        color = AMBER if busy else GREEN
        self.status_label.configure(text=f"● {text}", text_color=color)

    # ====================================================================
    # DASHBOARD TAB
    # ====================================================================
    def _build_dashboard_tab(self, tab):
        stats_row = ctk.CTkFrame(tab, fg_color="transparent")
        stats_row.pack(fill="x", padx=16, pady=(16, 12))
        stats_row.columnconfigure((0, 1, 2), weight=1)

        self.stat_total = self._stat_card(stats_row, "TOTAL REGISTERED", TEAL)
        self.stat_total[0].grid(row=0, column=0, sticky="ew", padx=(0, 8))
        self.stat_present = self._stat_card(stats_row, "PRESENT TODAY", GREEN)
        self.stat_present[0].grid(row=0, column=1, sticky="ew", padx=8)
        self.stat_late = self._stat_card(stats_row, "LATE TODAY", AMBER)
        self.stat_late[0].grid(row=0, column=2, sticky="ew", padx=(8, 0))

        ctk.CTkLabel(tab, text="Today's ledger", text_color=TEXT_MUTED,
                     font=("Segoe UI", 11, "bold")).pack(anchor="w", padx=16, pady=(4, 4))

        self.dash_list = ctk.CTkScrollableFrame(tab, fg_color=BG_APP, corner_radius=8)
        self.dash_list.pack(fill="both", expand=True, padx=16, pady=(0, 16))

    def _stat_card(self, parent, label, color):
        card = ctk.CTkFrame(parent, fg_color=BG_APP, corner_radius=10, border_color=BORDER, border_width=1)
        val = ctk.CTkLabel(card, text="0", text_color=color, font=("Consolas", 22, "bold"))
        val.pack(anchor="w", padx=14, pady=(10, 0))
        ctk.CTkLabel(card, text=label, text_color=TEXT_MUTED, font=("Segoe UI", 9, "bold")).pack(
            anchor="w", padx=14, pady=(0, 10))
        return card, val

    def _render_ledger(self, container, rows, search_text=""):
        for w in container.winfo_children():
            w.destroy()
        ft = (search_text or "").strip().lower()
        for (sid, name, date_, tstamp, status) in rows:
            if ft and ft not in name.lower() and ft not in sid.lower():
                continue
            row = ctk.CTkFrame(container, fg_color="transparent")
            row.pack(fill="x", pady=2)
            ctk.CTkLabel(row, text=f"#{sid}", text_color=TEXT_MUTED, font=FONT_MONO, width=60,
                         anchor="w").pack(side="left")
            ctk.CTkLabel(row, text=name, text_color=TEXT_LIGHT, anchor="w").pack(
                side="left", fill="x", expand=True)
            ctk.CTkLabel(row, text=tstamp, text_color=TEXT_MUTED, font=FONT_MONO, width=80).pack(side="left")
            color = AMBER if status == "Late" else GREEN
            ctk.CTkLabel(row, text=status.upper(), text_color=color, font=("Segoe UI", 9, "bold"),
                         width=70).pack(side="left")

    # ====================================================================
    # TAKE ATTENDANCE TAB
    # ====================================================================
    def _build_attendance_tab(self, tab):
        top = ctk.CTkFrame(tab, fg_color="transparent")
        top.pack(fill="x", padx=16, pady=16)

        self.track_btn = ctk.CTkButton(top, text="Take Attendance", fg_color=TEAL, text_color="#ffffff",
                                        hover_color="#4a3fd6", font=("Segoe UI", 13, "bold"), height=42,
                                        command=self.on_track_images)
        self.track_btn.pack(side="left")

        self.att_search_var = tk.StringVar()
        search = ctk.CTkEntry(top, textvariable=self.att_search_var, placeholder_text="Search by name or ID",
                               fg_color=BG_APP, border_color=BORDER, width=260)
        search.pack(side="right")
        self.att_search_var.trace_add("write", lambda *_: self.refresh_ledgers())

        self.att_list = ctk.CTkScrollableFrame(tab, fg_color=BG_APP, corner_radius=8)
        self.att_list.pack(fill="both", expand=True, padx=16, pady=(0, 16))

    def on_track_images(self):
        ok, reason = cascade_is_valid(HAARCASCADE_PATH)
        if not ok:
            self._msg("Cascade File Problem",
                       f"{reason}\n\nDownload the raw file from:\n{HAARCASCADE_RAW_URL}\n\n"
                       f"and save it as:\n{HAARCASCADE_PATH}")
            return
        if not busy_lock.acquire(blocking=False):
            self._msg("Busy", "Another camera operation is already running. Please wait.")
            return
        if not hasattr(cv2, "face"):
            self._msg("Missing Dependency", "Install opencv-contrib-python:\npip install opencv-contrib-python")
            busy_lock.release()
            return
        if not os.path.isfile(TRAINER_FILE):
            self._msg("Data Missing", "Please go to Register and click Save Profile to train the model first.")
            busy_lock.release()
            return
        if not os.path.isfile(STUDENT_DETAILS_XLSX):
            self._msg("Details Missing", "Student details are missing, please check.")
            busy_lock.release()
            return

        self.set_status("Opening camera…", busy=True)
        threading.Thread(target=self._track_images_worker, daemon=True).start()

    def _track_images_worker(self):
        cam = None
        try:
            recognizer = cv2.face.LBPHFaceRecognizer_create()
            recognizer.read(TRAINER_FILE)
            _next_serial, student_lookup = read_student_details()
            already_marked_ids = {r[0] for r in attendance_rows_cache}

            cam = open_camera()
            if not cam.isOpened():
                ui(self._msg, "Camera Error", "Could not open the webcam. Make sure it's not in use elsewhere.")
                return
            if not warm_up_camera(cam):
                ui(self._msg, "Camera Error", "The webcam opened but never delivered a frame.")
                return

            ui(self.set_status, "Scanning for faces…", True)
            faceCascade = cv2.CascadeClassifier(HAARCASCADE_PATH)
            seen_serials = set()
            attendance_rows = []
            frame_count = 0
            last_faces = []
            date = datetime.datetime.now().strftime('%d-%m-%Y')
            threshold = SETTINGS.get("confidence_threshold", 50)

            try:
                while True:
                    ret, im = cam.read()
                    if not ret or im is None or im.size == 0:
                        break
                    im = cv2.resize(im, (CAM_WIDTH, CAM_HEIGHT))
                    frame_count += 1
                    if frame_count % DETECT_EVERY_N_FRAMES == 0:
                        gray = cv2.cvtColor(im, cv2.COLOR_BGR2GRAY)
                        last_faces = faceCascade.detectMultiScale(gray, 1.2, 5, minSize=(60, 60))
                        for (x, y, w, h) in last_faces:
                            serial, conf = recognizer.predict(gray[y:y + h, x:x + w])
                            if conf < threshold and serial in student_lookup:
                                student_id, student_name = student_lookup[serial]
                                already_done = (student_id in already_marked_ids) or (serial in seen_serials)
                                if not already_done:
                                    seen_serials.add(serial)
                                    now_dt = datetime.datetime.now()
                                    now_str = now_dt.strftime('%H:%M:%S')
                                    status = compute_status(now_dt)
                                    attendance_rows.append([student_id, student_name, date, now_str, status])
                                    ui(self._add_attendance_record, student_id, student_name, date, now_str, status)
                    for (x, y, w, h) in last_faces:
                        cv2.rectangle(im, (x, y), (x + w, y + h), (66, 135, 245), 2)
                    cv2.imshow('Taking Attendance (press q to stop)', im)
                    if cv2.waitKey(1) & 0xFF == ord('q'):
                        break
            finally:
                cam.release()
                cv2.destroyAllWindows()

            if attendance_rows:
                append_attendance_rows(date, attendance_rows)
        except Exception as exc:
            ui(self._msg, "Attendance Error", f"Something went wrong:\n{exc}")
        finally:
            if cam is not None:
                try:
                    cam.release()
                except Exception:
                    pass
            ui(self.set_status, "Ready", False)
            busy_lock.release()

    def _add_attendance_record(self, sid, name, date_, tstamp, status):
        attendance_rows_cache.insert(0, (sid, name, date_, tstamp, status))
        self.refresh_ledgers()
        self.refresh_stats()
        beep()

    def refresh_ledgers(self):
        self._render_ledger(self.dash_list, attendance_rows_cache)
        self._render_ledger(self.att_list, attendance_rows_cache, self.att_search_var.get())

    # ====================================================================
    # STUDENTS TAB
    # ====================================================================
    def _build_students_tab(self, tab):
        top = ctk.CTkFrame(tab, fg_color="transparent")
        top.pack(fill="x", padx=16, pady=16)

        self.stu_search_var = tk.StringVar()
        search = ctk.CTkEntry(top, textvariable=self.stu_search_var, placeholder_text="Search students",
                               fg_color=BG_APP, border_color=BORDER, width=260)
        search.pack(side="left")
        self.stu_search_var.trace_add("write", lambda *_: self.refresh_students_list())

        ctk.CTkButton(top, text="Refresh", fg_color="#3a4152", hover_color="#4a5265",
                      command=self.refresh_students_list, width=100).pack(side="right")

        self.stu_list = ctk.CTkScrollableFrame(tab, fg_color=BG_APP, corner_radius=8)
        self.stu_list.pack(fill="both", expand=True, padx=16, pady=(0, 16))

    def refresh_students_list(self):
        for w in self.stu_list.winfo_children():
            w.destroy()
        _next_serial, students = read_student_details()
        ft = self.stu_search_var.get().strip().lower()

        if not students:
            ctk.CTkLabel(self.stu_list, text="No students registered yet — use the Register tab.",
                         text_color=TEXT_MUTED).pack(anchor="w", pady=20, padx=8)
            return

        for serial, (sid, name) in sorted(students.items(), key=lambda kv: kv[1][1].lower()):
            if ft and ft not in name.lower() and ft not in sid.lower():
                continue
            row = ctk.CTkFrame(self.stu_list, fg_color=BG_PANEL, corner_radius=8)
            row.pack(fill="x", pady=4)

            initials = "".join(p[0] for p in name.split()[:2]).upper() or "?"
            ctk.CTkLabel(row, text=initials, width=34, height=34, corner_radius=17, fg_color=GOLD,
                         text_color="#ffffff", font=("Segoe UI", 11, "bold")).pack(side="left", padx=10, pady=8)

            info = ctk.CTkFrame(row, fg_color="transparent")
            info.pack(side="left", fill="x", expand=True)
            ctk.CTkLabel(info, text=name, text_color=TEXT_LIGHT, font=("Segoe UI", 12, "bold"),
                         anchor="w").pack(anchor="w")
            ctk.CTkLabel(info, text=f"ID {sid}", text_color=TEXT_MUTED, font=FONT_MONO, anchor="w").pack(anchor="w")

            count = photo_count_for_serial(serial)
            count_color = AMBER if count < MIN_SAMPLES_WARN else TEXT_MUTED
            ctk.CTkLabel(row, text=f"{count} photos", text_color=count_color, font=("Segoe UI", 10)).pack(
                side="left", padx=12)

            ctk.CTkButton(row, text="Delete", fg_color="transparent", border_color=RED, border_width=1,
                          text_color=RED, hover_color="#2a1414", width=80, height=28,
                          command=lambda s=serial, n=name: self.delete_student(s, n)).pack(side="right", padx=10)

    def delete_student(self, serial, name):
        win = self._confirm_dialog(
            "Delete Student",
            f"Remove {name} and delete all their training photos?\n"
            f"You'll need to click Save Profile again afterwards to retrain the model.",
            on_yes=lambda: self._do_delete_student(serial, name),
        )

    def _do_delete_student(self, serial, name):
        remove_student_detail(serial)
        removed = delete_training_images(serial)
        self._msg("Student Removed", f"{name} was removed ({removed} training photo(s) deleted).\n"
                                      f"Go to Register and click Save Profile to retrain the model.")
        self.refresh_students_list()
        self.refresh_stats()

    # ====================================================================
    # REGISTER TAB
    # ====================================================================
    def _build_register_tab(self, tab):
        wrap = ctk.CTkFrame(tab, fg_color="transparent")
        wrap.pack(fill="both", expand=True, padx=16, pady=16)

        ctk.CTkLabel(wrap, text="New registration", text_color=TEXT_LIGHT,
                     font=("Segoe UI", 15, "bold")).pack(anchor="w")
        ctk.CTkLabel(wrap, text="Capture photos, then train the model.", text_color=TEXT_MUTED).pack(
            anchor="w", pady=(0, 16))

        ctk.CTkLabel(wrap, text="ID (numbers only)", text_color=TEXT_LIGHT).pack(anchor="w")
        vcmd_id = (self.register(self._validate_id_input), '%P')
        self.reg_id_entry = ctk.CTkEntry(wrap, fg_color=BG_APP, border_color=BORDER, width=320,
                                          validate="key", validatecommand=vcmd_id)
        self.reg_id_entry.pack(anchor="w", pady=(4, 14))

        ctk.CTkLabel(wrap, text="Name (letters and spaces only)", text_color=TEXT_LIGHT).pack(anchor="w")
        vcmd_name = (self.register(self._validate_name_input), '%P')
        self.reg_name_entry = ctk.CTkEntry(wrap, fg_color=BG_APP, border_color=BORDER, width=320,
                                            validate="key", validatecommand=vcmd_name)
        self.reg_name_entry.pack(anchor="w", pady=(4, 18))

        self.reg_message = ctk.CTkLabel(wrap, text="1) Take Images  →  2) Save Profile", text_color=TEXT_MUTED)
        self.reg_message.pack(anchor="w", pady=(0, 14))

        ctk.CTkButton(wrap, text="Take Images", fg_color=TEAL, text_color="#ffffff", hover_color="#4a3fd6",
                      width=320, height=42, command=self.on_take_images).pack(anchor="w", pady=(0, 10))
        ctk.CTkButton(wrap, text="Save Profile (train model)", fg_color=GOLD, text_color="#ffffff",
                      hover_color="#5a4dd9", width=320, height=42, command=self.on_save_profile).pack(anchor="w")

    @staticmethod
    def _validate_id_input(new_value):
        return new_value == "" or new_value.isdigit()

    @staticmethod
    def _validate_name_input(new_value):
        return new_value == "" or all(c.isalpha() or c.isspace() for c in new_value)

    def on_take_images(self):
        ok, reason = cascade_is_valid(HAARCASCADE_PATH)
        if not ok:
            self._msg("Cascade File Problem",
                       f"{reason}\n\nDownload the raw file from:\n{HAARCASCADE_RAW_URL}\n\n"
                       f"and save it as:\n{HAARCASCADE_PATH}")
            return
        if not busy_lock.acquire(blocking=False):
            self._msg("Busy", "Another camera operation is already running. Please wait.")
            return

        Id = self.reg_id_entry.get().strip()
        name = self.reg_name_entry.get().strip()
        if not Id:
            self.reg_message.configure(text="Enter an ID (numbers only)", text_color=RED)
            busy_lock.release()
            return
        if not name:
            self.reg_message.configure(text="Enter a name (letters only)", text_color=RED)
            busy_lock.release()
            return

        self.set_status("Opening camera…", busy=True)
        self.reg_message.configure(text="Opening camera... click CAPTURE (or press C) to save each photo",
                                    text_color=TEXT_MUTED)
        threading.Thread(target=self._take_images_worker, args=(Id, name), daemon=True).start()

    def _take_images_worker(self, Id, name):
        cam = None
        try:
            serial, _existing = read_student_details()
            cam = open_camera()
            if not cam.isOpened():
                ui(self._msg, "Camera Error", "Could not open the webcam.")
                return
            if not warm_up_camera(cam):
                ui(self._msg, "Camera Error", "The webcam opened but never delivered a frame.")
                return

            ui(self.set_status, "Camera ready — click CAPTURE to save a photo", True)
            detector = cv2.CascadeClassifier(HAARCASCADE_PATH)
            sampleNum = 0
            frame_count = 0
            last_faces = []

            win_name = 'Take Images  —  click CAPTURE or press C to save, Q to finish'
            cv2.namedWindow(win_name)
            btn_w, btn_h, btn_margin = 190, 54, 20
            btn_x1 = CAM_WIDTH - btn_w - btn_margin
            btn_y1 = CAM_HEIGHT - btn_h - btn_margin
            btn_x2 = CAM_WIDTH - btn_margin
            btn_y2 = CAM_HEIGHT - btn_margin
            capture_event = {"go": False}
            last_capture_time = {"t": 0.0}
            CAPTURE_COOLDOWN = 0.35
            flash_until = 0.0

            def _on_mouse(event, x, y, flags, param):
                if event == cv2.EVENT_LBUTTONDOWN and btn_x1 <= x <= btn_x2 and btn_y1 <= y <= btn_y2:
                    capture_event["go"] = True

            cv2.setMouseCallback(win_name, _on_mouse)

            try:
                while True:
                    ret, img = cam.read()
                    if not ret:
                        break
                    frame_count += 1
                    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
                    if frame_count % DETECT_EVERY_N_FRAMES == 0:
                        last_faces = detector.detectMultiScale(gray, 1.3, 5, minSize=(60, 60))
                    face_box = max(last_faces, key=lambda f: f[2] * f[3]) if len(last_faces) > 0 else None
                    if face_box is not None:
                        fx, fy, fw, fh = face_box
                        cv2.rectangle(img, (fx, fy), (fx + fw, fy + fh), (66, 135, 245), 2)

                    key = cv2.waitKey(1) & 0xFF
                    now_t = time.time()
                    wants_capture = (capture_event["go"] or key == ord('c')) and \
                                     (now_t - last_capture_time["t"] > CAPTURE_COOLDOWN)
                    capture_event["go"] = False

                    if wants_capture:
                        last_capture_time["t"] = now_t
                        if face_box is not None:
                            fx, fy, fw, fh = face_box
                            sampleNum += 1
                            filename = f"{name}.{serial}.{Id}.{sampleNum}.jpg"
                            cv2.imwrite(os.path.join(TRAINING_IMAGE_DIR, filename), gray[fy:fy + fh, fx:fx + fw])
                            flash_until = now_t + 0.15

                    btn_color = (34, 197, 94) if now_t < flash_until else (245, 130, 40)
                    cv2.rectangle(img, (btn_x1, btn_y1), (btn_x2, btn_y2), btn_color, -1)
                    cv2.putText(img, "CAPTURE (C)", (btn_x1 + 18, btn_y1 + 34),
                                cv2.FONT_HERSHEY_SIMPLEX, 0.65, (255, 255, 255), 2)
                    cv2.putText(img, f"Captured: {sampleNum}/{MAX_SAMPLES}", (10, 30),
                                cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 255, 0), 2)
                    if face_box is None:
                        cv2.putText(img, "No face detected", (10, 58),
                                    cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 165, 255), 2)
                    cv2.imshow(win_name, img)
                    if key == ord('q'):
                        break
                    if sampleNum >= MAX_SAMPLES:
                        break
            finally:
                cam.release()
                cv2.destroyAllWindows()

            if sampleNum == 0:
                ui(self.reg_message.configure, {"text": "No photos captured — please try again", "text_color": RED})
                return

            append_student_detail(serial, Id, name)
            ui(self.reg_message.configure, {"text": f"Saved {sampleNum} photo(s) for ID {Id}", "text_color": GREEN})
            ui(self.refresh_stats)
            ui(self.refresh_students_list)
        except Exception as exc:
            ui(self._msg, "Camera Error", f"Something went wrong while capturing images:\n{exc}")
        finally:
            if cam is not None:
                try:
                    cam.release()
                except Exception:
                    pass
            ui(self.set_status, "Ready", False)
            busy_lock.release()

    def on_save_profile(self):
        assure_path_exists(TRAINING_LABEL_DIR)
        if not os.path.isfile(PASSWORD_FILE):
            self._prompt_dialog("Set Password",
                                 "No password is set yet. This is ONE shared password used every time\n"
                                 "you click Save Profile. Enter a new password:", show="*",
                                 on_submit=self._set_new_password)
            return
        self._prompt_dialog("Password Required", "Enter your Save Profile password:", show="*",
                             on_submit=self._check_password_and_train)

    def _set_new_password(self, value):
        if not value or not value.strip():
            self._msg("No Password Entered", "Password not set. Please try again.")
            return
        with open(PASSWORD_FILE, "w") as f:
            f.write(hash_password(value.strip()))
        self._msg("Password Registered", "Password set. Click Save Profile again to train the model.")

    def _check_password_and_train(self, value):
        if value is None:
            return
        with open(PASSWORD_FILE, "r") as f:
            key = f.read().strip()
        if hash_password(value.strip()) != key:
            self._msg("Wrong Password", "You entered the wrong password.")
            return
        self._train_images()

    def _train_images(self):
        if not busy_lock.acquire(blocking=False):
            self._msg("Busy", "Another camera operation is already running. Please wait.")
            return
        _faces_unused, ids = getImagesAndLabels(TRAINING_IMAGE_DIR)
        if not ids:
            busy_lock.release()
            self._msg("No Registrations", "Please register someone first with Take Images.")
            return

        counts = Counter(ids)
        _next_serial, students = read_student_details()
        low = [(serial, students.get(serial, ('?', 'Unknown'))[1], c)
               for serial, c in counts.items() if c < MIN_SAMPLES_WARN]

        def _proceed():
            self.set_status("Training model…", busy=True)
            self.reg_message.configure(text="Training in progress...", text_color=TEXT_MUTED)
            threading.Thread(target=self._train_images_worker, daemon=True).start()

        if low:
            names_list = "\n".join(f"  • {name} — only {c} photo(s)" for _serial, name, c in low)
            self._confirm_dialog(
                "Low Photo Count",
                f"These students have fewer than {MIN_SAMPLES_WARN} training photos, "
                f"which may make recognition unreliable:\n\n{names_list}\n\nTrain anyway?",
                on_yes=_proceed, on_no=busy_lock.release,
            )
        else:
            _proceed()

    def _train_images_worker(self):
        try:
            if not hasattr(cv2, "face"):
                ui(self._msg, "Missing Dependency",
                   "cv2.face is not available.\nInstall opencv-contrib-python.")
                return
            recognizer = cv2.face.LBPHFaceRecognizer_create()
            faces, ids = getImagesAndLabels(TRAINING_IMAGE_DIR)
            if not faces:
                ui(self._msg, "No Registrations", "Please register someone first with Take Images.")
                return
            recognizer.train(faces, np.array(ids))
            recognizer.save(TRAINER_FILE)
            ui(self.reg_message.configure, {"text": "Profile saved successfully", "text_color": GREEN})
            ui(self.refresh_stats)
        except Exception as exc:
            ui(self._msg, "Training Error", f"Something went wrong while training:\n{exc}")
        finally:
            ui(self.set_status, "Ready", False)
            busy_lock.release()

    # ====================================================================
    # REPORTS TAB
    # ====================================================================
    def _build_reports_tab(self, tab):
        top = ctk.CTkFrame(tab, fg_color="transparent")
        top.pack(fill="x", padx=16, pady=16)

        ctk.CTkLabel(top, text="From (dd-mm-yyyy)", text_color=TEXT_MUTED).grid(row=0, column=0, sticky="w")
        ctk.CTkLabel(top, text="To (dd-mm-yyyy)", text_color=TEXT_MUTED).grid(row=0, column=1, sticky="w", padx=(12, 0))

        today = datetime.date.today()
        week_ago = today - datetime.timedelta(days=7)
        self.rep_from_entry = ctk.CTkEntry(top, fg_color=BG_APP, border_color=BORDER, width=140)
        self.rep_from_entry.insert(0, week_ago.strftime("%d-%m-%Y"))
        self.rep_from_entry.grid(row=1, column=0, sticky="w", pady=(2, 0))

        self.rep_to_entry = ctk.CTkEntry(top, fg_color=BG_APP, border_color=BORDER, width=140)
        self.rep_to_entry.insert(0, today.strftime("%d-%m-%Y"))
        self.rep_to_entry.grid(row=1, column=1, sticky="w", padx=(12, 0), pady=(2, 0))

        ctk.CTkButton(top, text="Run report", fg_color=TEAL, text_color="#ffffff", hover_color="#4a3fd6",
                      command=self.run_report).grid(row=1, column=2, padx=(16, 0))
        ctk.CTkButton(top, text="Export CSV", fg_color="#3a4152", hover_color="#4a5265",
                      command=self.export_report_csv).grid(row=1, column=3, padx=(8, 0))

        self.rep_summary = ctk.CTkLabel(tab, text="Pick a date range and click Run report.", text_color=TEXT_MUTED)
        self.rep_summary.pack(anchor="w", padx=16, pady=(4, 8))

        self.rep_list = ctk.CTkScrollableFrame(tab, fg_color=BG_APP, corner_radius=8)
        self.rep_list.pack(fill="both", expand=True, padx=16, pady=(0, 16))
        self._last_report_rows = []

    def _parse_date(self, text):
        return datetime.datetime.strptime(text.strip(), "%d-%m-%Y").date()

    def run_report(self):
        try:
            start = self._parse_date(self.rep_from_entry.get())
            end = self._parse_date(self.rep_to_entry.get())
        except ValueError:
            self._msg("Invalid Date", "Use dd-mm-yyyy for both dates, e.g. 01-08-2026.")
            return
        if start > end:
            self._msg("Invalid Range", "The From date must be before the To date.")
            return
        if (end - start).days > 92:
            self._msg("Range Too Large", "Please pick a range of 92 days or fewer.")
            return

        records = load_attendance_range(start, end)
        total_days = (end - start).days + 1
        per_student = defaultdict(lambda: {"name": "", "present": 0, "late": 0})
        for sid, name, _date, _time, status in records:
            per_student[sid]["name"] = name
            per_student[sid]["present"] += 1
            if status == "Late":
                per_student[sid]["late"] += 1

        rows = []
        for sid, d in sorted(per_student.items(), key=lambda kv: -kv[1]["present"]):
            pct = round(100 * d["present"] / total_days) if total_days else 0
            rows.append((sid, d["name"], d["present"], d["late"], pct))
        self._last_report_rows = rows

        for w in self.rep_list.winfo_children():
            w.destroy()

        if not rows:
            ctk.CTkLabel(self.rep_list, text="No attendance recorded in this range.",
                         text_color=TEXT_MUTED).pack(anchor="w", pady=20, padx=8)
        else:
            header = ctk.CTkFrame(self.rep_list, fg_color="transparent")
            header.pack(fill="x", pady=(0, 4))
            for txt, w in (("ID", 60), ("Name", 200), ("Days present", 100), ("Late", 60), ("Attendance %", 100)):
                ctk.CTkLabel(header, text=txt, text_color=TEXT_MUTED, font=("Segoe UI", 9, "bold"),
                             width=w, anchor="w").pack(side="left")
            for sid, name, present, late, pct in rows:
                row = ctk.CTkFrame(self.rep_list, fg_color="transparent")
                row.pack(fill="x", pady=2)
                ctk.CTkLabel(row, text=sid, text_color=TEXT_MUTED, font=FONT_MONO, width=60, anchor="w").pack(side="left")
                ctk.CTkLabel(row, text=name, text_color=TEXT_LIGHT, width=200, anchor="w").pack(side="left")
                ctk.CTkLabel(row, text=str(present), text_color=TEXT_LIGHT, width=100, anchor="w").pack(side="left")
                ctk.CTkLabel(row, text=str(late), text_color=AMBER, width=60, anchor="w").pack(side="left")
                pct_color = GREEN if pct >= 90 else (AMBER if pct >= 75 else RED)
                ctk.CTkLabel(row, text=f"{pct}%", text_color=pct_color, width=100, anchor="w").pack(side="left")

        self.rep_summary.configure(
            text=f"{len(rows)} student(s) with at least one mark, over {total_days} day(s) "
                 f"({start.strftime('%d-%m-%Y')} to {end.strftime('%d-%m-%Y')})."
        )

    def export_report_csv(self):
        if not self._last_report_rows:
            self._msg("Nothing To Export", "Run a report first.")
            return
        out_path = os.path.join(ATTENDANCE_DIR, f"Report_{int(time.time())}.csv")
        import csv
        with open(out_path, "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["Id", "Name", "Days Present", "Late Count", "Attendance %"])
            writer.writerows(self._last_report_rows)
        self._msg("Exported", f"Saved to:\n{out_path}")

    # ====================================================================
    # SETTINGS TAB
    # ====================================================================
    def _build_settings_tab(self, tab):
        wrap = ctk.CTkFrame(tab, fg_color="transparent")
        wrap.pack(fill="both", expand=True, padx=16, pady=16)

        ctk.CTkLabel(wrap, text="Class start time (HH:MM, 24h)", text_color=TEXT_LIGHT).pack(anchor="w")
        self.set_time_entry = ctk.CTkEntry(wrap, fg_color=BG_APP, border_color=BORDER, width=120)
        self.set_time_entry.insert(0, SETTINGS["class_start_time"])
        self.set_time_entry.pack(anchor="w", pady=(4, 14))

        ctk.CTkLabel(wrap, text="Recognition confidence threshold (lower = stricter)",
                     text_color=TEXT_LIGHT).pack(anchor="w")
        self.set_conf_slider = ctk.CTkSlider(wrap, from_=20, to=90, number_of_steps=70, width=320)
        self.set_conf_slider.set(SETTINGS["confidence_threshold"])
        self.set_conf_slider.pack(anchor="w", pady=(4, 2))
        self.set_conf_value = ctk.CTkLabel(wrap, text=str(SETTINGS["confidence_threshold"]), text_color=TEXT_MUTED)
        self.set_conf_value.pack(anchor="w", pady=(0, 14))
        self.set_conf_slider.configure(command=lambda v: self.set_conf_value.configure(text=str(int(v))))

        ctk.CTkLabel(wrap, text="Camera index (0 = default webcam)", text_color=TEXT_LIGHT).pack(anchor="w")
        self.set_cam_entry = ctk.CTkEntry(wrap, fg_color=BG_APP, border_color=BORDER, width=80)
        self.set_cam_entry.insert(0, str(SETTINGS["camera_index"]))
        self.set_cam_entry.pack(anchor="w", pady=(4, 14))

        self.set_sound_var = tk.BooleanVar(value=SETTINGS["sound_on"])
        ctk.CTkCheckBox(wrap, text="Play a sound on each successful recognition",
                         variable=self.set_sound_var, text_color=TEXT_LIGHT,
                         fg_color=TEAL, hover_color="#4a3fd6").pack(anchor="w", pady=(0, 20))

        ctk.CTkButton(wrap, text="Save settings", fg_color=GOLD, text_color="#ffffff", hover_color="#5a4dd9",
                      width=200, command=self.save_settings_clicked).pack(anchor="w")

        ctk.CTkButton(wrap, text="Open Attendance Folder", fg_color="#3a4152", hover_color="#4a5265",
                      width=200, command=self.open_attendance_folder).pack(anchor="w", pady=(24, 0))

    def save_settings_clicked(self):
        global SETTINGS
        time_txt = self.set_time_entry.get().strip()
        try:
            h, m = time_txt.split(":")
            datetime.time(int(h), int(m))
        except Exception:
            self._msg("Invalid Time", "Use HH:MM 24-hour format, e.g. 09:00.")
            return
        try:
            cam_idx = int(self.set_cam_entry.get().strip())
        except ValueError:
            self._msg("Invalid Camera Index", "Camera index must be a whole number, e.g. 0 or 1.")
            return

        SETTINGS = {
            "class_start_time": time_txt,
            "confidence_threshold": int(self.set_conf_slider.get()),
            "camera_index": cam_idx,
            "sound_on": bool(self.set_sound_var.get()),
        }
        save_settings(SETTINGS)
        self._msg("Settings Saved", "New settings will apply to the next camera session.")
        self.refresh_ledgers()

    def open_attendance_folder(self):
        try:
            if platform.system() == "Windows":
                os.startfile(ATTENDANCE_DIR)  # noqa
            elif platform.system() == "Darwin":
                subprocess.Popen(["open", ATTENDANCE_DIR])
            else:
                subprocess.Popen(["xdg-open", ATTENDANCE_DIR])
        except Exception as exc:
            self._msg("Could Not Open Folder", f"{exc}\n\nFolder path:\n{ATTENDANCE_DIR}")

    # ====================================================================
    # Shared refresh / dialogs
    # ====================================================================
    def refresh_stats(self):
        _next_serial, students = read_student_details()
        total_reg = len(students)
        present = len(attendance_rows_cache)
        late = sum(1 for r in attendance_rows_cache if r[4] == 'Late')
        self.stat_total[1].configure(text=str(total_reg))
        self.stat_present[1].configure(text=str(present))
        self.stat_late[1].configure(text=str(late))

    def refresh_all(self):
        global attendance_rows_cache, today_date_str
        attendance_rows_cache, today_date_str = load_today_attendance()
        self.refresh_ledgers()
        self.refresh_stats()
        self.refresh_students_list()

    def _msg(self, title, message):
        win = ctk.CTkToplevel(self)
        win.title(title)
        win.configure(fg_color=BG_PANEL)
        win.geometry("420x220")
        win.transient(self)
        win.grab_set()
        ctk.CTkLabel(win, text=title, text_color=TEXT_LIGHT, font=("Segoe UI", 14, "bold")).pack(
            anchor="w", padx=20, pady=(20, 6))
        ctk.CTkLabel(win, text=message, text_color=TEXT_MUTED, wraplength=380, justify="left").pack(
            anchor="w", padx=20)
        ctk.CTkButton(win, text="OK", fg_color=TEAL, text_color="#ffffff", command=win.destroy, width=100).pack(
            anchor="e", padx=20, pady=20, side="bottom")

    def _confirm_dialog(self, title, message, on_yes=None, on_no=None):
        win = ctk.CTkToplevel(self)
        win.title(title)
        win.configure(fg_color=BG_PANEL)
        win.geometry("440x240")
        win.transient(self)
        win.grab_set()
        ctk.CTkLabel(win, text=title, text_color=TEXT_LIGHT, font=("Segoe UI", 14, "bold")).pack(
            anchor="w", padx=20, pady=(20, 6))
        ctk.CTkLabel(win, text=message, text_color=TEXT_MUTED, wraplength=400, justify="left").pack(
            anchor="w", padx=20)

        btn_row = ctk.CTkFrame(win, fg_color="transparent")
        btn_row.pack(anchor="e", padx=20, pady=20, side="bottom")

        def _yes():
            win.destroy()
            if on_yes:
                on_yes()

        def _no():
            win.destroy()
            if on_no:
                on_no()

        ctk.CTkButton(btn_row, text="Cancel", fg_color="#3a4152", command=_no, width=100).pack(side="right", padx=(10, 0))
        ctk.CTkButton(btn_row, text="Yes, continue", fg_color=RED, command=_yes, width=140).pack(side="right")

    def _prompt_dialog(self, title, message, show=None, on_submit=None):
        win = ctk.CTkToplevel(self)
        win.title(title)
        win.configure(fg_color=BG_PANEL)
        win.geometry("420x220")
        win.transient(self)
        win.grab_set()
        ctk.CTkLabel(win, text=message, text_color=TEXT_LIGHT, wraplength=380, justify="left").pack(
            anchor="w", padx=20, pady=(20, 10))
        entry = ctk.CTkEntry(win, fg_color=BG_APP, border_color=BORDER, show=show, width=340)
        entry.pack(padx=20)
        entry.focus_set()

        def _submit(_e=None):
            val = entry.get()
            win.destroy()
            if on_submit:
                on_submit(val)

        entry.bind("<Return>", _submit)
        ctk.CTkButton(win, text="OK", fg_color=TEAL, text_color="#ffffff", command=_submit, width=100).pack(
            anchor="e", padx=20, pady=20, side="bottom")


def getImagesAndLabels(path):
    imagePaths = [os.path.join(path, f) for f in os.listdir(path) if f.lower().endswith(('.jpg', '.jpeg', '.png'))]
    faces, Ids = [], []
    for imagePath in imagePaths:
        try:
            pilImage = Image.open(imagePath).convert('L')
            imageNp = np.array(pilImage, 'uint8')
            serial = int(os.path.split(imagePath)[-1].split(".")[1])
            faces.append(imageNp)
            Ids.append(serial)
        except (IndexError, ValueError):
            continue
    return faces, Ids


if __name__ == "__main__":
    app = FaceTrackApp()
    app.mainloop()